"""XLSX transfer for private dictionaries and existing AI memory."""
import json
import re
import secrets
from datetime import date, datetime, time
from io import BytesIO
from urllib.parse import quote
from zipfile import ZipFile

from fastapi import Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field
from starlette.concurrency import run_in_threadpool
from backend.document_fields import key
from backend.saved_data import DataInput

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 1000
MAX_TEXT = 2 * 1024 * 1024
AI_FIELDS = ['Pole', 'Wartość']

def export_book(fields, rows):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Dane'
    for values in [fields, *[[row.get(field, '') for field in fields] for row in rows]]:
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            # All dictionary values are text, including leading zeros and strings beginning with '='.
            cell.data_type = 's'
            cell.number_format = '@'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='6654D9')
        sheet.column_dimensions[cell.column_letter].width = min(45, max(20, len(cell.value) + 5))
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()

def read_book(data):
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024 or len(archive.infolist()) > 2000:
                raise HTTPException(413, 'Arkusz jest zbyt duży po rozpakowaniu.')
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = book.worksheets[0]
            sheet.reset_dimensions()
            rows = sheet.iter_rows(max_col=61)
            header = next(rows)
            while header and header[-1].value is None: header = header[:-1]
            if not header: raise HTTPException(422, 'Pierwszy wiersz musi zawierać nazwy kolumn.')
            if len(header) > 60: raise HTTPException(422, 'Arkusz może zawierać maksymalnie 60 kolumn.')
            fields = []
            for cell in header:
                if cell.data_type == 'f' or not isinstance(cell.value, str) or not cell.value.strip():
                    raise HTTPException(422, 'Każda kolumna musi mieć tekstową nazwę w pierwszym wierszu.')
                fields.append(cell.value.strip())
            if len({key(field) for field in fields}) != len(fields): raise HTTPException(422, 'Nazwy kolumn nie mogą się powtarzać.')
            result, size = [], 0
            for number, row in enumerate(rows, 2):
                if number > MAX_ROWS + 1:
                    raise HTTPException(413, 'Import obsługuje maksymalnie 1000 wierszy.')
                if any(cell.value is not None for cell in row[len(fields):]):
                    raise HTTPException(422, f'Wiersz {number} zawiera dane w kolumnie bez nagłówka.')
                values = {}
                for field, cell in zip(fields, row):
                    value = cell.value
                    if cell.data_type in ('f', 'e'):
                        raise HTTPException(422, f'Wiersz {number}: zamień formuły i błędy Excela na zwykłe wartości.')
                    if value is None: value = ''
                    elif isinstance(value, (datetime, date, time)): value = value.isoformat()
                    elif isinstance(value, bool): value = 'Tak' if value else 'Nie'
                    elif isinstance(value, int) and re.fullmatch('0+', cell.number_format or ''):
                        value = str(value).zfill(len(cell.number_format))
                    else: value = str(value)
                    values[field] = value.strip()
                if not any(values.values()): continue
                if len(result) >= MAX_ROWS: raise HTTPException(413, 'Import obsługuje maksymalnie 1000 wierszy.')
                size += sum(len(value) for value in values.values())
                if size > MAX_TEXT: raise HTTPException(413, 'Za dużo tekstu w arkuszu.')
                result.append(values)
            if not result: raise HTTPException(422, 'Arkusz nie zawiera wpisów do importu.')
            return fields, result, sheet.title
        finally: book.close()
    except HTTPException: raise
    except Exception as exc:
        raise HTTPException(422, 'Nie można odczytać pliku. Wybierz poprawny, niezabezpieczony plik .xlsx.') from exc

class ImportInput(BaseModel):
    version: int | None = Field(default=None, strict=True, ge=1)
    rows: list[dict[str, str]] = Field(min_length=1, max_length=MAX_ROWS)

def install(app, db, user, now):
    def resource(con, identifier, current):
        if identifier == 'ai': return {'id': 'ai', 'name': 'AI', 'fields': AI_FIELDS, 'version': None}
        row = con.execute('SELECT * FROM dictionaries WHERE id=? AND user_id=?', (identifier, current['id'])).fetchone()
        if not row: raise HTTPException(404, 'Nie znaleziono słownika.')
        return {**dict(row), 'fields': json.loads(row['fields'])}

    def normalized(rows, fields, ai=False):
        output, size = [], 0
        targets = {key(field): field for field in fields}
        for number, row in enumerate(rows, 2):
            if len(row) > 60: raise HTTPException(422, 'Za dużo kolumn.')
            keys = [key(field) for field in row]
            if len(set(keys)) != len(keys) or set(keys) - set(targets):
                raise HTTPException(422, 'Kolumny Excela nie pasują do pól słownika. Najpierw dodaj brakujące pola lub popraw nagłówki.')
            values = dict.fromkeys(fields, '')
            for label, value in row.items():
                if len(value) > 2000: raise HTTPException(422, f'Wiersz {number}: wartość jest za długa (maksymalnie 2000 znaków).')
                values[targets[key(label)]] = value.strip()
            if not any(values.values()): raise HTTPException(422, f'Wiersz {number} jest pusty.')
            if ai:
                try: DataInput(name=values['Pole'][:100], fields={values['Pole']: values['Wartość']})
                except ValueError as exc: raise HTTPException(422, f'Wiersz {number}: uzupełnij poprawne pole i wartość dla AI.') from exc
            size += sum(len(value) for value in values.values())
            if size > MAX_TEXT: raise HTTPException(413, 'Za dużo tekstu do importu.')
            output.append(values)
        return output

    @app.get('/api/dictionaries/{identifier}/excel')
    def export(identifier: str, current=Depends(user)):
        with db() as con:
            item = resource(con, identifier, current)
            if identifier == 'ai':
                rows = [{'Pole': label, 'Wartość': value} for row in con.execute('SELECT fields FROM saved_document_data WHERE user_id=? ORDER BY updated_at DESC,id', (current['id'],)) for label, value in json.loads(row['fields']).items()]
            else:
                rows = [json.loads(row['values_json']) for row in con.execute('SELECT values_json FROM dictionary_entries WHERE dictionary_id=? ORDER BY updated_at DESC,id', (identifier,))]
        filename = quote(item['name'] + '.xlsx', safe='')
        return Response(export_book(item['fields'], rows), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                        headers={'Content-Disposition': f"attachment; filename=dictionary.xlsx; filename*=UTF-8''{filename}"})

    @app.post('/api/dictionaries/{identifier}/excel/preview')
    async def preview(identifier: str, file: UploadFile = File(...), current=Depends(user)):
        with db() as con: item = resource(con, identifier, current)
        try:
            if not (file.filename or '').lower().endswith('.xlsx'): raise HTTPException(422, 'Wybierz plik Excel w formacie .xlsx.')
            data = await file.read(MAX_BYTES + 1)
            if len(data) > MAX_BYTES: raise HTTPException(413, 'Plik może mieć maksymalnie 10 MB.')
            fields, rows, sheet = await run_in_threadpool(read_book, data)
            if set(key(field) for field in fields) - set(key(field) for field in item['fields']):
                raise HTTPException(422, 'Nagłówki Excela nie pasują do słownika. Najpierw dodaj brakujące pola lub popraw nagłówki.')
            rows = normalized(rows, item['fields'], identifier == 'ai')
            return {'fields': item['fields'], 'rows': rows, 'count': len(rows), 'version': item['version'], 'sheet': sheet}
        finally: await file.close()

    @app.post('/api/dictionaries/{identifier}/excel/import', status_code=201)
    def commit(identifier: str, body: ImportInput, current=Depends(user)):
        with db() as con:
            con.begin_write()
            item = resource(con, identifier, current)
            if identifier != 'ai' and item['version'] != body.version:
                raise HTTPException(409, 'Pola słownika zmieniły się. Wczytaj plik ponownie.')
            rows = normalized(body.rows, item['fields'], identifier == 'ai')
            for values in rows:
                identifier_new = secrets.token_hex(16)
                if identifier == 'ai':
                    fields = {key(values['Pole']): values['Wartość']}
                    con.execute('INSERT INTO saved_document_data(id,user_id,name,fields,updated_at) VALUES(?,?,?,?,?)',
                        (identifier_new, current['id'], values['Pole'][:100], json.dumps(fields, ensure_ascii=False), now()))
                else:
                    con.execute('INSERT INTO dictionary_entries(id,dictionary_id,values_json,updated_at) VALUES(?,?,?,?)',
                        (identifier_new, identifier, json.dumps(values, ensure_ascii=False), now()))
            return {'imported': len(rows)}
