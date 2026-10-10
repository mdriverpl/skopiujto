from io import BytesIO
import pytest
from openpyxl import Workbook, load_workbook
from backend.test_storage import app_client
from backend.test_access import ADMIN, READER, create_reader, login
from backend.dictionary_excel import export_book, read_book

def excel(fields, rows):
    book = Workbook(); sheet = book.active; sheet.append(fields)
    for row in rows: sheet.append(row)
    result = BytesIO(); book.save(result); book.close()
    return result.getvalue()

def dictionary(client):
    return client.post('/api/dictionaries', json={'name': 'Pracownicy', 'fields': ['Imię', 'Telefon', 'E-mail']}).json()

def test_excel_export_and_preview_roundtrip_preserve_text_and_do_not_write_until_commit(app_client):
    _, client = app_client
    item = dictionary(client); path = '/api/dictionaries/' + item['id']
    values = {'Imię': '=SUM(1,2)', 'Telefon': '00123', 'E-mail': 'a@example.com'}
    client.post(path + '/entries', json={'values': values})
    response = client.get(path + '/excel')
    assert response.status_code == 200
    assert 'spreadsheetml' in response.headers['content-type']
    book = load_workbook(BytesIO(response.content))
    assert book.active['A2'].data_type == 's' and book.active['A2'].value == '=SUM(1,2)'
    assert book.active['B2'].value == '00123' and book.active.freeze_panes == 'A2'
    book.close()
    preview = client.post(path + '/excel/preview', files={'file': ('data.xlsx', response.content)}).json()
    assert preview['rows'] == [values] and preview['count'] == 1
    assert client.get(path + '/entries').json()['dictionary']['count'] == 1
    imported = client.post(path + '/excel/import', json={'version': preview['version'], 'rows': preview['rows']})
    assert imported.status_code == 201 and imported.json()['imported'] == 1
    assert client.get(path + '/entries').json()['dictionary']['count'] == 2

def test_import_matches_headers_by_name_ignores_blank_rows_and_fills_missing_columns(app_client):
    _, client = app_client
    item = dictionary(client); path = '/api/dictionaries/' + item['id']
    file = excel(['telefon', ' imię '], [['00123', 'Anna'], [None, None], ['456', 'Jan']])
    response = client.post(path + '/excel/preview', files={'file': ('data.xlsx', file)})
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview['count'] == 2 and preview['rows'][0] == {'Imię': 'Anna', 'Telefon': '00123', 'E-mail': ''}
    assert client.post(path + '/excel/import', json={'version': 1, 'rows': preview['rows']}).status_code == 201

def test_excel_import_checks_schema_version_ownership_and_atomic_validation(app_client):
    _, client = app_client
    create_reader(client)
    item = dictionary(client); path = '/api/dictionaries/' + item['id']
    rows = [{'Imię': 'Anna'}, {'Unknown': 'Bad'}]
    assert client.post(path + '/excel/import', json={'version': 1, 'rows': rows}).status_code == 422
    assert client.get(path + '/entries').json()['entries'] == []
    assert client.post(path + '/excel/import', json={'version': 2, 'rows': [{'Imię': 'Anna'}]}).status_code == 409
    file = excel(['Imię'], [['Anna']])
    login(client, READER)
    assert client.get(path + '/excel').status_code == 404
    assert client.post(path + '/excel/preview', files={'file': ('data.xlsx', file)}).status_code == 404
    assert client.post(path + '/excel/import', json={'version': 1, 'rows': [{'Imię': 'Anna'}]}).status_code == 404
    login(client, ADMIN); client.post('/api/logout')
    assert client.get(path + '/excel').status_code == 401

@pytest.mark.parametrize('file', [b'bad file', excel(['Imię', ' imię '], [['A', 'B']]),
    excel(['Unknown'], [['A']]), excel(['Imię'], [['=SUM(1,2)']]), excel(['Imię'], [])])
def test_invalid_workbook_does_not_change_entries(app_client, file):
    _, client = app_client
    item = dictionary(client); path = '/api/dictionaries/' + item['id']
    assert client.post(path + '/excel/preview', files={'file': ('bad.xlsx', file)}).status_code == 422
    assert client.get(path + '/entries').json()['entries'] == []

def test_ai_excel_transfer_uses_existing_private_memory(app_client):
    _, client = app_client
    client.post('/api/saved-data/remember', json={'name': 'City', 'fields': {'Miasto': 'Warszawa'}, 'source_key': 'one'})
    exported = client.get('/api/dictionaries/ai/excel')
    fields, rows, _ = read_book(exported.content)
    assert fields == ['Pole', 'Wartość'] and rows == [{'Pole': 'miasto', 'Wartość': 'Warszawa'}]
    file = export_book(fields, [{'Pole': 'Adres', 'Wartość': 'Długa 1'}])
    preview = client.post('/api/dictionaries/ai/excel/preview', files={'file': ('AI.xlsx', file)}).json()
    assert client.post('/api/dictionaries/ai/excel/import', json={'rows': preview['rows']}).status_code == 201
    assert len(client.get('/api/saved-data').json()) == 2
    assert client.post('/api/dictionaries/ai/excel/import', json={'rows': [{'Pole': 'Adres', 'Wartość': ''}]}).status_code == 422

def test_excel_limits_and_numeric_zero_format(app_client):
    _, client = app_client
    item = dictionary(client); path = '/api/dictionaries/' + item['id']
    assert client.post(path + '/excel/preview', files={'file': ('old.xls', b'bad')}).status_code == 422
    assert client.post(path + '/excel/preview', files={'file': ('huge.xlsx', b'x' * (10 * 1024 * 1024 + 1))}).status_code == 413
    book = Workbook(); book.active.append(['Telefon']); book.active.append([123]); book.active['A2'].number_format = '00000'
    data = BytesIO(); book.save(data); book.close()
    assert read_book(data.getvalue())[1] == [{'Telefon': '00123'}]
    too_many = excel(['Imię'], [['Anna']] + [[None]] * 1000 + [['Jan']])
    assert client.post(path + '/excel/preview', files={'file': ('many.xlsx', too_many)}).status_code == 413
